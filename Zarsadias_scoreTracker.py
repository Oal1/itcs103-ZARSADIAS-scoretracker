#filename: surname_scoreTracker.py

import tkinter as tk
from tkinter import ttk
from tkinter import messagebox
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
import os

EXCEL_FILE = "student_scores.xlsx"

# Initialize workbook
def excel_file():
    if not os.path.exists(EXCEL_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = "Scores"

        headers = ["Student Name", "Score", "Remarks"]
        ws.append(headers)

        # Make headers bold and center-aligned
        for col_num, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        wb.save(EXCEL_FILE)


# Function to calculate and write average row
def update_average_row(ws):
    scores = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if row[1].value is not None and isinstance(row[1].value, (int, float)):
            scores.append(row[1].value)

    if not scores:
        return

    avg_score = sum(scores) / len(scores)
    remark = "Passed" if avg_score >= 75 else "Failed"

    avg_row_index = len(scores) + 2
    ws.cell(row=avg_row_index, column=1).value = "AVERAGE"
    ws.cell(row=avg_row_index, column=2).value = avg_score
    ws.cell(row=avg_row_index, column=3).value = remark

    # Remove any extra rows after the average (if user deleted a record)
    if ws.max_row > avg_row_index:
        for row in ws.iter_rows(min_row=avg_row_index + 1, max_row=ws.max_row):
            for cell in row:
                cell.value = None

# Save student data
def save_data():
    name = name_entry.get().strip()
    score_input = score_entry.get().strip()

    if not name or not score_input:
        messagebox.showwarning("Missing Input", "Please fill in both the student name and score.")
        return

    if any(char.isdigit() for char in name):
        messagebox.showerror("Invalid Name", "Student name must not contain numbers.")
        return

    try:
        score = float(score_input)
        
        # Check if the score is within the valid range (0 to 100)
        if score < 0 or score > 100:
            messagebox.showerror("Invalid Score", "Score must be between 0 and 100.")
            return
    except ValueError:
        messagebox.showerror("Invalid Input", "Score must be a valid number.")
        return

    remarks = "Passed" if score >= 75 else "Failed"

    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    # Remove previous average row (if it exists)
    if ws.max_row >= 3:
        last_row = ws.max_row
        if ws.cell(row=last_row, column=1).value == "AVERAGE":
            ws.delete_rows(last_row)

    ws.append([name, score, remarks])
    update_average_row(ws)

    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        ws.column_dimensions[column_letter].width = max_length + 2  # +2 for padding

    wb.save(EXCEL_FILE)

    messagebox.showinfo("Success", f"Record saved for {name}")
    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)



# Display records in a new window
def display_records():
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    display_window = tk.Toplevel(window)
    display_window.title("All Records")
    display_window.geometry("400x400")

    tree = ttk.Treeview(display_window, columns=("Name", "Score", "Remarks"), show="headings", height=20)
    tree.pack(fill=tk.BOTH, expand=True)

    tree.heading("Name", text="Student Name")
    tree.heading("Score", text="Score")
    tree.heading("Remarks", text="Remarks")

    tree.column("Name", anchor=tk.W, width=150)
    tree.column("Score", anchor=tk.CENTER, width=80)
    tree.column("Remarks", anchor=tk.CENTER, width=100)

    # Add rows from Excel
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(cell is not None for cell in row):  # skip completely empty rows
            tree.insert("", tk.END, values=row)

# GUI
window = tk.Tk()
window.title("Score Tracker")

tk.Label(window, text="Student Name:").grid(row=0, column=0, padx=10, pady=5)
name_entry = tk.Entry(window)
name_entry.grid(row=0, column=1, padx=10, pady=5)

tk.Label(window, text="Score:").grid(row=1, column=0, padx=10, pady=5)
score_entry = tk.Entry(window)
score_entry.grid(row=1, column=1, padx=10, pady=5)

tk.Button(window, text="Save Record", command=save_data).grid(row=2, column=0, columnspan=2, pady=10)
tk.Button(window, text="Display All Records", command=display_records).grid(row=3, column=0, columnspan=2, pady=5)

excel_file()
window.mainloop()